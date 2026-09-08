#!/usr/bin/env python3
"""Client for the public UFIP / Énergies et Mobilités custom-value export.

The A4C Rotterdam reference is explicitly the UFIP page's Rotterdam quotation in EUR/litre,
source Thomson-Reuters, published as a 5-day moving average. The downloader validates that
public source contract before accepting an export so a future unit/method change cannot be
silently consumed under the same column name.
"""
from __future__ import annotations

import io
import math
import re
from datetime import date

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

UFIP_CUSTOM_URL = "https://valeurs.ufip.fr/datas/custom"
USER_AGENT = "A4C-observatoires/2.0 (+public-data research)"
BROWSER_USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/152.0.0.0 Safari/537.36"
)
GAZOLE_HEADER_PREFIX = "GAZOLE (Rotterdam)"
ROTTERDAM_UNIT = "EUR/L"
ROTTERDAM_REFERENCE_SOURCE = "Thomson-Reuters"
ROTTERDAM_SMOOTHING = "5-day moving average"
ROTTERDAM_MIN_EUR_L = 0.05
ROTTERDAM_MAX_EUR_L = 5.0


def _format_date(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def validate_ufip_source_contract(raw_html: str) -> None:
    """Require the public UFIP page to still describe the expected Rotterdam series."""
    text = BeautifulSoup(raw_html, "html.parser").get_text(" ", strip=True)
    text = " ".join(text.replace("\xa0", " ").split())
    folded = text.casefold().replace("–", "-").replace("—", "-")
    if "cotations rotterdam" not in folded:
        raise RuntimeError("UFIP source contract changed: Rotterdam quotation section not found")
    unit_ok = any(token in folded for token in ("€ /litre", "€/litre", "€ / litre", "eur/litre", "eur / litre"))
    if not unit_ok:
        raise RuntimeError("UFIP source contract changed: Rotterdam unit is no longer EUR/litre")
    if "thomson-reuters" not in folded and "thomson reuters" not in folded:
        raise RuntimeError("UFIP source contract changed: Thomson-Reuters source label not found")
    if not re.search(r"moyennes?\s+mobiles?\s+sur\s+5\s+jours", folded):
        raise RuntimeError("UFIP source contract changed: 5-day moving-average label not found")


def _fetch_validated_contract_page(session, *, timeout: int = 90):
    """Fetch the documented UFIP form and retry once if a minimal edge-page variant is served.

    GitHub-hosted runners can receive a stripped/minimal first response from the UFIP front
    door. We still fail closed on semantics: the retry must explicitly contain the Rotterdam
    EUR/litre, Thomson-Reuters and 5-day moving-average wording before any export is accepted.
    """
    first = session.get(UFIP_CUSTOM_URL, timeout=timeout)
    first.raise_for_status()
    try:
        validate_ufip_source_contract(first.text)
        return first
    except RuntimeError as first_error:
        retry = session.get(
            UFIP_CUSTOM_URL,
            timeout=timeout,
            headers={
                "User-Agent": BROWSER_USER_AGENT,
                "Accept-Language": "fr-FR,fr;q=0.9,en;q=0.8",
                "Cache-Control": "no-cache",
                "Pragma": "no-cache",
            },
        )
        retry.raise_for_status()
        try:
            validate_ufip_source_contract(retry.text)
        except RuntimeError as retry_error:
            raise RuntimeError(
                "UFIP source contract unavailable after browser-header retry; "
                f"first={first_error}; retry={retry_error}"
            ) from retry_error
        return retry


def _validated_rotterdam_value(raw_value, *, context: str) -> float:
    try:
        value = float(raw_value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid Rotterdam EUR/L value in {context}: {raw_value!r}") from exc
    if not math.isfinite(value) or not ROTTERDAM_MIN_EUR_L <= value <= ROTTERDAM_MAX_EUR_L:
        raise ValueError(f"Implausible Rotterdam EUR/L value in {context}: {raw_value!r}")
    return value


def parse_rotterdam_gazole_xlsx(raw: bytes) -> pd.DataFrame:
    """Parse the two-column UFIP export into date/value rows."""
    if not raw.startswith(b"PK"):
        raise ValueError("UFIP response is not an XLSX ZIP container")
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    if not wb.worksheets:
        raise ValueError("UFIP workbook has no worksheet")
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("UFIP workbook is empty")
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    try:
        date_col = header.index("Date")
    except ValueError as exc:
        raise ValueError(f"UFIP workbook has no Date column: {header}") from exc
    fuel_col = next((i for i, value in enumerate(header) if value.startswith(GAZOLE_HEADER_PREFIX)), None)
    if fuel_col is None:
        raise ValueError(f"UFIP workbook has no Rotterdam Gazole column: {header}")

    parsed = []
    for row_number, row in enumerate(rows[1:], start=2):
        if date_col >= len(row) or fuel_col >= len(row):
            continue
        raw_date, raw_value = row[date_col], row[fuel_col]
        if raw_date is None or raw_value is None:
            continue
        if hasattr(raw_date, "date"):
            d = raw_date.date()
        elif isinstance(raw_date, date):
            d = raw_date
        else:
            d = pd.to_datetime(raw_date, dayfirst=True).date()
        value = _validated_rotterdam_value(raw_value, context=f"XLSX row {row_number} ({d})")
        parsed.append((d, value))
    df = pd.DataFrame(parsed, columns=["date", "rotterdam_eur_l"])
    if not df.empty:
        df = df.sort_values("date").drop_duplicates("date", keep="last").reset_index(drop=True)
    return df


def fetch_rotterdam_gazole(
    start_date: date,
    end_date: date,
    *,
    session: requests.Session | None = None,
    timeout: int = 90,
) -> pd.DataFrame:
    """Download the validated UFIP Rotterdam Gazole series for a custom period."""
    if end_date < start_date:
        raise ValueError("end_date must be >= start_date")
    own_session = session is None
    s = session or requests.Session()
    s.headers.setdefault("User-Agent", USER_AGENT)
    try:
        first = _fetch_validated_contract_page(s, timeout=timeout)
        soup = BeautifulSoup(first.text, "html.parser")
        token = soup.select_one('input[name="ufp_token"]')
        if token is None or not token.get("value"):
            raise RuntimeError("UFIP ufp_token not found in validated custom export form")
        payload = {
            "ufp_token": token["value"],
            "day_from": _format_date(start_date),
            "day_to": _format_date(end_date),
            "cotations[gazole]": "on",
        }
        response = s.post(UFIP_CUSTOM_URL, data=payload, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
        frame = parse_rotterdam_gazole_xlsx(response.content)
        if not frame.empty:
            outside = frame[(frame["date"] < start_date) | (frame["date"] > end_date)]
            if not outside.empty:
                raise RuntimeError("UFIP export returned Rotterdam dates outside the requested interval")
        return frame
    finally:
        if own_session:
            s.close()


def expand_daily(observations: pd.DataFrame, start_date: date, end_date: date) -> pd.DataFrame:
    """Forward-fill weekends/holidays from the last UFIP observation."""
    if end_date < start_date:
        raise ValueError("end_date must be >= start_date")
    calendar = pd.DataFrame({"date": pd.date_range(start_date, end_date, freq="D").date})
    source = observations.copy()
    if not source.empty:
        source["date"] = pd.to_datetime(source["date"]).dt.date
        source = source.sort_values("date").drop_duplicates("date", keep="last")
        source["rotterdam_observed"] = True
    merged = calendar.merge(source, on="date", how="left")
    if "rotterdam_observed" not in merged.columns:
        merged["rotterdam_observed"] = False
    else:
        merged["rotterdam_observed"] = merged["rotterdam_observed"].eq(True)
    merged["rotterdam_eur_l"] = pd.to_numeric(merged.get("rotterdam_eur_l"), errors="coerce").ffill()
    merged["rotterdam_carried"] = merged["rotterdam_eur_l"].notna() & ~merged["rotterdam_observed"]
    return merged
